"""A4: the `state` filter repeats, on the list and on the XLSX export (real Postgres).

`?state=a&state=b` lists the applications in state `a` or `b`. A single value keeps
working, and without `state` the list is not filtered.
"""

from __future__ import annotations

import io
from collections.abc import AsyncIterator

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.modules.auth.principal import Principal
from tests._support.read_models import (
    as_principal,
    build_read_api,
    create_app,
    fire,
    seed_read_world,
    staff,
)

pytestmark = pytest.mark.integration

_READER = Principal(
    sub="reader", permissions={"application.read", "application.export"}
)


@pytest.fixture
async def maker(
    migrated: tuple[str, str], engine: Engine
) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    eng = create_async_engine(migrated[1])
    yield async_sessionmaker(eng, expire_on_commit=False)
    await eng.dispose()


async def test_state_filter_takes_several_states(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed = await seed_read_world(maker)
    in_eingang = await create_app(maker, seed, actor=seed.owner_sub, title="Eins")
    in_pruefung = await create_app(maker, seed, actor=seed.owner_sub, title="Zwei")
    in_genehmigt = await create_app(maker, seed, actor=seed.owner_sub, title="Drei")
    await fire(maker, in_pruefung, seed.to_review_id, staff(seed))
    await fire(maker, in_genehmigt, seed.to_review_id, staff(seed))
    await fire(maker, in_genehmigt, seed.approve_id, staff(seed))

    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, _READER)
    with TestClient(api) as client:
        both = client.get(
            "/api/applications",
            params=[("state", str(seed.eingang_id)), ("state", str(seed.pruefung_id))],
        )
        one = client.get("/api/applications", params={"state": str(seed.genehmigt_id)})
        unfiltered = client.get("/api/applications")
        bad = client.get("/api/applications", params=[("state", "kein-uuid")])
        export = client.get(
            "/api/applications/export.xlsx",
            params=[("state", str(seed.eingang_id)), ("state", str(seed.genehmigt_id))],
        )

    assert both.status_code == 200, both.text
    assert {i["id"] for i in both.json()["items"]} == {str(in_eingang), str(in_pruefung)}
    assert both.json()["total"] == 2
    assert [i["id"] for i in one.json()["items"]] == [str(in_genehmigt)]
    assert unfiltered.json()["total"] == 3
    assert bad.status_code == 422

    assert export.status_code == 200, export.text
    sheet = load_workbook(io.BytesIO(export.content)).active
    assert sheet is not None
    titles = {str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert titles == {"Eins", "Drei"}
