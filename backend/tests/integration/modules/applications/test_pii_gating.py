"""O21: the `isPII` fields reach only a reader with the PII right (real Postgres).

The right is `application.read`, `application.read_all` or admin. The applicant
(magic link) and the logged-in creator read the own data. A member who reads only
through the Gremium read scope gets the detail and the version history without the
`isPII` fields and without the applicant block. An editor with `application.manage`
alone cannot see them either, so the patch keeps their stored values. The XLSX export
holds no form field value at all. The search finds no application by a hidden value.
The attachments of an `isPII` field stay out of the list and the downloads too.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import Engine, select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.modules.applications.models import Application
from app.modules.applications.service import ApplicationsService
from app.modules.auth.principal import Principal
from app.modules.files.models import Attachment
from app.modules.forms.models import FormField, FormVersion
from tests._support.read_models import (
    IBAN,
    MEMBER_NAME,
    OWNER_NAME,
    ReadSeed,
    as_applicant,
    as_principal,
    build_read_api,
    create_app,
    get_json,
    patch,
    seed_read_world,
    staff,
)

pytestmark = pytest.mark.integration

_NEW_IBAN = "DE89370400440532013000"


@pytest.fixture
async def maker(
    migrated: tuple[str, str], engine: Engine
) -> AsyncIterator[async_sessionmaker[AsyncSession]]:
    eng = create_async_engine(migrated[1])
    yield async_sessionmaker(eng, expire_on_commit=False)
    await eng.dispose()


async def _world(
    maker: async_sessionmaker[AsyncSession],
) -> tuple[ReadSeed, uuid.UUID]:
    """An application by the creator, with a second version that changes the IBAN."""
    seed = await seed_read_world(maker)
    app_id = await create_app(maker, seed, actor=seed.owner_sub)
    await patch(
        maker,
        app_id,
        {"title": "Antrag", "iban": _NEW_IBAN, "note": "zweite"},
        changed_by=seed.owner_sub,
    )
    return seed, app_id


def _detail(api: FastAPI, app_id: uuid.UUID) -> dict:
    body = get_json(api, f"/api/applications/{app_id}")
    assert isinstance(body, dict)
    return body


def _versions(api: FastAPI, app_id: uuid.UUID) -> list:
    body = get_json(api, f"/api/applications/{app_id}/versions")
    assert isinstance(body, list)
    return body


@pytest.mark.parametrize(
    "principal",
    [
        Principal(sub="reader", permissions={"application.read"}),
        Principal(sub="global", permissions={"application.read_all"}),
        Principal(sub="root", roles=["admin"]),
    ],
    ids=["read", "read_all", "admin"],
)
async def test_reader_with_the_right_sees_pii(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
    principal: Principal,
) -> None:
    _, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, principal)

    detail = _detail(api, app_id)
    assert detail["data"]["iban"] == _NEW_IBAN
    assert detail["applicant"]["email"] == "antrag@example.org"
    assert detail["hiddenKeys"] == []

    versions = _versions(api, app_id)
    assert versions[0]["data"]["iban"] == IBAN
    assert versions[1]["diff"]["changed"]["iban"] == {"old": IBAN, "new": _NEW_IBAN}
    assert versions[1]["changedKeys"] == ["iban", "note"]


async def test_committee_reader_gets_no_pii(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    # Gremium member without a global permission: reads through the view cost centre.
    as_principal(api, Principal(sub=seed.member_sub, display_name=MEMBER_NAME))

    detail = _detail(api, app_id)
    assert "iban" not in detail["data"]
    assert detail["data"]["note"] == "zweite"
    assert detail["applicant"] is None
    # The client hides exactly these fields in the edit form.
    assert detail["hiddenKeys"] == ["iban"]

    versions = _versions(api, app_id)
    assert [v["version"] for v in versions] == [1, 2]
    assert all("iban" not in v["data"] for v in versions)
    assert "iban" not in versions[1]["diff"]["changed"]
    assert versions[1]["diff"]["changed"]["note"] == {"old": "erste Notiz", "new": "zweite"}
    assert versions[1]["changedKeys"] == ["note"]
    assert versions[1]["changedBy"] == OWNER_NAME
    assert IBAN not in str(versions) and _NEW_IBAN not in str(versions)


async def test_applicant_and_creator_read_the_own_pii(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)

    as_applicant(api, app_id)
    assert _detail(api, app_id)["data"]["iban"] == _NEW_IBAN

    as_principal(api, Principal(sub=seed.owner_sub, display_name=OWNER_NAME))
    detail = _detail(api, app_id)
    assert detail["data"]["iban"] == _NEW_IBAN
    assert detail["applicant"]["name"] == "Anna Antrag"
    assert detail["isOwner"] is True
    assert detail["hiddenKeys"] == []


async def test_reader_with_the_right_gets_an_unanswered_pii_field_as_editable(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A missing key is not a hidden key: the full reader can still fill the field."""
    seed = await seed_read_world(maker)
    app_id = await create_app(maker, seed, actor=seed.owner_sub)
    await patch(maker, app_id, {"title": "Antrag", "note": "ohne"}, changed_by=seed.owner_sub)
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, Principal(sub="reader", permissions={"application.read"}))

    detail = _detail(api, app_id)
    assert "iban" not in detail["data"]
    assert detail["hiddenKeys"] == []


async def test_editor_without_read_keeps_the_stored_pii(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    # `application.manage` without `application.read`: the editor reads no IBAN and
    # sends the form back without it, or with a value it never saw.
    as_principal(api, staff(seed))
    with TestClient(api) as client:
        without = client.patch(
            f"/api/applications/{app_id}",
            json={"data": {"title": "Antrag", "note": "dritte"}},
        )
        forged = client.patch(
            f"/api/applications/{app_id}",
            json={"data": {"title": "Antrag", "iban": "DE00", "note": "vierte"}},
        )
    assert without.status_code == 200, without.text
    assert "iban" not in without.json()["data"]
    assert forged.status_code == 200, forged.text
    assert "iban" not in forged.json()["data"]

    async with maker() as session:
        row = await session.get(Application, app_id)
        assert row is not None
        assert row.data["iban"] == _NEW_IBAN
        assert row.data["note"] == "vierte"


async def test_archive_response_strips_pii_without_read(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, Principal(sub="archivar", permissions={"application.archive"}))
    with TestClient(api) as client:
        archived = client.post(f"/api/applications/{app_id}/archive")
        restored = client.delete(f"/api/applications/{app_id}/archive")
    assert archived.status_code == 200, archived.text
    assert "iban" not in archived.json()["data"]
    assert "iban" not in restored.json()["data"]


async def test_idempotent_archive_strips_pii_without_read(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The no-op path of archive and unarchive applies O21 too."""
    _, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, Principal(sub="archivar", permissions={"application.archive"}))
    with TestClient(api) as client:
        unarchived_noop = client.delete(f"/api/applications/{app_id}/archive")
        first = client.post(f"/api/applications/{app_id}/archive")
        second = client.post(f"/api/applications/{app_id}/archive")
    for resp in (unarchived_noop, first, second):
        assert resp.status_code == 200, resp.text
        assert "iban" not in resp.json()["data"]
        assert resp.json()["hiddenKeys"] == ["iban"]
    assert second.json()["archivedAt"] == first.json()["archivedAt"]


def _search_total(api: FastAPI, q: str) -> int:
    body = get_json(api, f"/api/applications?q={q}")
    assert isinstance(body, dict)
    return int(body["total"])


async def test_search_finds_no_application_by_a_hidden_pii_value(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The search is no oracle for a PII value (O21).

    A Gremium reader without the PII right finds the application by its title, but
    not by its IBAN. The reader with the right and the creator find it by both.
    """
    seed, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)

    as_principal(api, Principal(sub=seed.member_sub, display_name=MEMBER_NAME))
    assert _search_total(api, "Antrag") == 1
    assert _search_total(api, _NEW_IBAN) == 0

    as_principal(api, Principal(sub="reader", permissions={"application.read"}))
    assert _search_total(api, _NEW_IBAN) == 1

    as_principal(api, Principal(sub=seed.owner_sub, display_name=OWNER_NAME))
    assert _search_total(api, _NEW_IBAN) == 1
    assert str(app_id) in str(get_json(api, f"/api/applications?q={_NEW_IBAN}"))


async def test_creator_without_read_gets_no_internal_comment(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The creator without a read permission reads as the applicant (A12)."""
    seed, app_id = await _world(maker)
    async with maker() as session:
        for visibility in ("internal", "public"):
            await ApplicationsService(session).add_comment(
                app_id,
                author=seed.member_sub,
                author_kind="principal",
                body=f"{visibility} note",
                visibility=visibility,
            )
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, Principal(sub=seed.owner_sub, display_name=OWNER_NAME))
    comments = get_json(api, f"/api/applications/{app_id}/comments")
    assert isinstance(comments, list)
    assert [c["body"] for c in comments] == ["public note"]
    with TestClient(api) as client:
        resp = client.post(
            f"/api/applications/{app_id}/comments",
            json={"body": "intern", "visibility": "internal"},
        )
    assert resp.status_code == 403, resp.text

    # A member with the read permission still reads both.
    as_principal(api, Principal(sub="reader", permissions={"application.read"}))
    comments = get_json(api, f"/api/applications/{app_id}/comments")
    assert isinstance(comments, list)
    assert sorted(c["body"] for c in comments) == ["internal note", "public note"]


async def test_editor_without_read_is_not_blocked_by_a_missing_required_pii_field(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A required isPII field without a stored value does not stop the blind editor."""
    seed = await seed_read_world(maker)
    app_id = await create_app(maker, seed, actor=seed.owner_sub)
    # The IBAN goes away, then the field becomes required: old data, new rule.
    await patch(maker, app_id, {"title": "Antrag", "note": "ohne"}, changed_by=seed.owner_sub)
    async with maker() as session:
        await session.execute(
            update(FormField)
            .where(
                FormField.key == "iban",
                FormField.form_version_id.in_(
                    select(FormVersion.id).where(FormVersion.application_type_id == seed.type_id)
                ),
            )
            .values(required=True)
        )
        await session.commit()

    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, staff(seed))
    with TestClient(api) as client:
        blind = client.patch(
            f"/api/applications/{app_id}", json={"data": {"title": "Antrag", "note": "neu"}}
        )
    assert blind.status_code == 200, blind.text

    # The reader who sees the field still gets the 422.
    as_principal(
        api, Principal(sub="editor", permissions={"application.read", "application.manage"})
    )
    with TestClient(api) as client:
        full = client.patch(
            f"/api/applications/{app_id}", json={"data": {"title": "Antrag", "note": "neu"}}
        )
    assert full.status_code == 422, full.text
    assert "iban" in full.text


async def test_xlsx_export_holds_no_pii(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, app_id = await _world(maker)
    api = build_read_api(migrated[1], monkeypatch)
    as_principal(api, Principal(sub="export", permissions={"application.export"}))
    with TestClient(api) as client:
        resp = client.get("/api/applications/export.xlsx")
    assert resp.status_code == 200, resp.text
    sheet = load_workbook(io.BytesIO(resp.content)).active
    assert sheet is not None
    cells = [str(c) for row in sheet.iter_rows(values_only=True) for c in row if c]
    assert "Antrag" in cells
    assert "Status seit" in cells
    assert not any(IBAN in c or _NEW_IBAN in c for c in cells)
    assert not any("antrag@example.org" in c for c in cells)
    assert str(app_id) not in cells


async def _attachments(
    maker: async_sessionmaker[AsyncSession], app_id: uuid.UUID
) -> tuple[uuid.UUID, uuid.UUID, uuid.UUID]:
    """Three quarantined attachments: by PII field key, by PII answer, and a plain one."""
    async with maker() as session:
        by_key = Attachment(
            application_id=app_id,
            field_key="iban",
            filename="ausweis.pdf",
            mime="application/pdf",
            size=10,
            storage_key="k1",
        )
        by_ref = Attachment(
            application_id=app_id,
            field_key=None,
            filename="nachweis.pdf",
            mime="application/pdf",
            size=10,
            storage_key="k2",
        )
        plain = Attachment(
            application_id=app_id,
            field_key="note",
            filename="angebot.pdf",
            mime="application/pdf",
            size=10,
            storage_key="k3",
        )
        session.add_all([by_key, by_ref, plain])
        await session.flush()
        # The answer of an isPII field refers to the attachment, like a file field.
        row = await session.get(Application, app_id)
        assert row is not None
        row.data = {**row.data, "iban": str(by_ref.id)}
        await session.commit()
        return by_key.id, by_ref.id, plain.id


def _attachment_codes(api: FastAPI, ids: tuple[uuid.UUID, ...]) -> list[tuple[int, int]]:
    """The status of the URL route and of the download route for each attachment."""
    with TestClient(api) as client:
        return [
            (
                client.get(f"/api/attachments/{i}").status_code,
                client.get(f"/api/attachments/{i}/download").status_code,
            )
            for i in ids
        ]


async def test_committee_reader_gets_no_pii_attachment(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seed, app_id = await _world(maker)
    by_key, by_ref, plain = await _attachments(maker, app_id)
    api = build_read_api(migrated[1], monkeypatch)

    as_principal(api, Principal(sub=seed.member_sub, display_name=MEMBER_NAME))
    listed = get_json(api, f"/api/applications/{app_id}/attachments")
    assert isinstance(listed, list)
    assert [a["id"] for a in listed] == [str(plain)]
    # 404 for the hidden files, the quarantine gate (409) for the plain one.
    assert _attachment_codes(api, (by_key, by_ref, plain)) == [
        (404, 404),
        (404, 404),
        (409, 409),
    ]


@pytest.mark.parametrize("who", ["reader", "creator", "applicant"])
async def test_pii_readers_get_every_attachment(
    migrated: tuple[str, str],
    maker: async_sessionmaker[AsyncSession],
    monkeypatch: pytest.MonkeyPatch,
    who: str,
) -> None:
    seed, app_id = await _world(maker)
    ids = await _attachments(maker, app_id)
    api = build_read_api(migrated[1], monkeypatch)
    if who == "reader":
        as_principal(api, Principal(sub="reader", permissions={"application.read"}))
    elif who == "creator":
        as_principal(api, Principal(sub=seed.owner_sub, display_name=OWNER_NAME))
    else:
        as_applicant(api, app_id)

    listed = get_json(api, f"/api/applications/{app_id}/attachments")
    assert isinstance(listed, list)
    assert sorted(a["id"] for a in listed) == sorted(str(i) for i in ids)
    assert _attachment_codes(api, ids) == [(409, 409)] * 3
