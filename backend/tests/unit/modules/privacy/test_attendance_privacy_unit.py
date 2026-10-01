"""Unit tests for the attendance note in the GDPR flows (Z2, A7).

The reason of an excuse (`meeting_attendance.note`) is personal data. The Art. 15
export lists the attendance of the principal with the reason. The Art. 17 erasure
of a principal clears the reason and keeps the status, because the protocols
carry the status.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql
from sqlalchemy.sql.dml import Update

from app.modules.auth.models import Principal
from app.modules.livevote.models import Meeting, MeetingAttendance
from app.modules.privacy.service import (
    AuskunftService,
    PrincipalService,
    build_auskunft_workbook,
)
from tests._support.privacy_fakes import FakeResult, FakeSession, result


class _RecordingSession(FakeSession):
    """Privacy fake session that also keeps every executed statement."""

    def __init__(self, **channels: Any) -> None:
        super().__init__(**channels)
        self.statements: list[Any] = []

    async def execute(self, stmt: Any) -> FakeResult:
        self.statements.append(stmt)
        return await super().execute(stmt)


def _principal() -> Principal:
    return Principal(
        id=uuid4(),
        sub="kc-1",
        email="m@example.org",
        display_name="Member",
        active=True,
    )


async def test_erase_clears_the_attendance_note() -> None:
    principal = _principal()
    db: Any = _RecordingSession(gets=[principal])
    await PrincipalService(db).erase(principal.id, actor="admin")
    updates = [
        s
        for s in db.statements
        if isinstance(s, Update)
        and getattr(s.table, "name", None) == MeetingAttendance.__tablename__
    ]
    assert len(updates) == 1
    sql = str(updates[0].compile(dialect=postgresql.dialect()))
    assert "SET note=" in sql
    # Only the rows with a reason change, so an older (self, absent) row is left alone.
    assert "note IS NOT NULL" in sql
    assert "principal_id" in sql


async def test_auskunft_lists_the_attendance_with_the_note() -> None:
    principal = _principal()
    meeting = Meeting(id=uuid4(), gremium_id=uuid4(), title="StuPa 3", date=date(2026, 9, 1))
    rec = MeetingAttendance(
        meeting_id=meeting.id,
        principal_id=principal.id,
        status="excused",
        source="self",
        note="Krank",
    )
    db: Any = _RecordingSession(
        scalars=[result()],  # no applicant
        scalar=[principal],
        execute=[result((rec, meeting))],
    )
    data = await AuskunftService(db).collect("m@example.org")
    assert data["attendance"] == [
        {
            "meetingId": meeting.id,
            "meetingTitle": "StuPa 3",
            "meetingDate": date(2026, 9, 1),
            "status": "excused",
            "source": "self",
            "note": "Krank",
        }
    ]


async def test_auskunft_without_principal_has_no_attendance() -> None:
    db: Any = _RecordingSession(scalars=[result()], scalar=[None])
    data = await AuskunftService(db).collect("nobody@example.org")
    assert data["attendance"] == []
    assert db.statements == []


def test_workbook_has_an_attendance_sheet() -> None:
    meeting_id = uuid4()
    raw = build_auskunft_workbook(
        email="m@example.org",
        applications=[],
        versions=[],
        principal=None,
        attendance=[
            {
                "meetingId": meeting_id,
                "meetingTitle": "StuPa 3",
                "meetingDate": date(2026, 9, 1),
                "status": "excused",
                "source": "self",
                # A reason is free text: the sheet neutralizes a formula.
                "note": "=HYPERLINK(1)",
            },
            {
                "meetingId": None,
                "meetingTitle": None,
                "meetingDate": None,
                "status": "present",
                "source": "lead",
                "note": None,
            },
        ],
    )
    ws = load_workbook(BytesIO(raw))["Anwesenheit"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Sitzungs-ID", "Sitzung", "Datum", "Status", "Erfasst durch", "Grund")
    assert rows[1][:5] == (str(meeting_id), "StuPa 3", "2026-09-01", "excused", "self")
    assert not str(rows[1][5]).startswith("=")
    assert rows[2][1:] == (None, None, "present", "lead", None)
