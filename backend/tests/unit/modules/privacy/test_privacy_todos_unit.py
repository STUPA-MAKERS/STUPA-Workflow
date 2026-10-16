"""The to-do part of the privacy module: the erasure and the Art. 15 export."""

from __future__ import annotations

import uuid
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.modules.auth.models import Principal
from app.modules.privacy import service as privacy_service
from app.modules.privacy.service import (
    AuskunftService,
    PrincipalService,
    build_auskunft_workbook,
)
from tests._support.privacy_fakes import fake_session, result

_AT = datetime(2026, 6, 16, 12, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def _quiet(monkeypatch: pytest.MonkeyPatch) -> None:
    async def _none(*_a: object, **_k: object) -> list[object]:
        return []

    async def _synced(*_a: object, **_k: object) -> bool:
        return True

    monkeypatch.setattr(privacy_service, "erase_candidacies", _none)
    monkeypatch.setattr(privacy_service, "erase_candidate_names", _none)
    monkeypatch.setattr(privacy_service, "sync_principal_memberships", _synced)


async def test_the_erasure_removes_the_todo_data(monkeypatch: pytest.MonkeyPatch) -> None:
    erased: list[uuid.UUID] = []

    async def _erase(_s: Any, principal_id: uuid.UUID) -> None:
        erased.append(principal_id)

    monkeypatch.setattr(privacy_service, "erase_todos", _erase)
    principal = Principal(id=uuid.uuid4(), sub="kc", email="u@x.org", active=True)
    await PrincipalService(fake_session(gets=[principal])).erase(
        principal.id, actor="admin", commit=False
    )
    assert erased == [principal.id]


async def test_the_export_asks_for_the_todos_of_the_account(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    asked: list[list[uuid.UUID]] = []

    async def _export(_s: Any, ids: list[uuid.UUID]) -> dict[str, Any]:
        asked.append(ids)
        return {"cards": [{"title": "A"}], "comments": [{"body": "B"}]}

    monkeypatch.setattr(privacy_service, "todo_export", _export)
    principal = Principal(id=uuid.uuid4(), sub="kc", email="x@example.org", active=True)
    db = fake_session(scalars=[result(), result()], scalar=[principal])
    data = await AuskunftService(db).collect("x@example.org")
    assert asked == [[principal.id]]
    assert data["todos"] == [{"title": "A"}] and data["todo_comments"] == [{"body": "B"}]
    anonymous = fake_session(scalars=[result(), result()], scalar=[None])
    await AuskunftService(anonymous).collect("y@example.org")
    assert asked[-1] == []


def test_the_workbook_has_the_todo_sheets() -> None:
    raw = build_auskunft_workbook(
        email="x@example.org",
        applications=[],
        versions=[],
        principal=None,
        todos=[
            {
                "board": "StuPa",
                "title": "=1+1",
                "column": "Offen",
                "dueDate": date(2026, 11, 1),
                "createdAt": _AT,
            },
            {"board": None, "title": None, "column": None, "dueDate": None, "createdAt": None},
        ],
        todo_comments=[
            {"card": "A", "body": "Hi", "createdAt": _AT, "deletedAt": _AT},
            {"card": None, "body": None, "createdAt": None, "deletedAt": None},
        ],
    )
    book = load_workbook(BytesIO(raw))
    cards = [[c.value for c in row] for row in book["To-Dos"].iter_rows()]
    assert cards[0] == ["Board", "Titel", "Spalte", "Fällig", "Erstellt"]
    assert cards[1][0] == "StuPa" and cards[1][3] == "2026-11-01"
    assert not str(cards[1][1]).startswith("=")
    assert cards[2] == [None, None, None, None, None]
    comments = [[c.value for c in row] for row in book["To-Do-Kommentare"].iter_rows()]
    assert comments[0] == ["To-Do", "Zeitpunkt", "Gelöscht", "Text"]
    assert comments[1] == ["A", "2026-06-16 12:00", "2026-06-16 12:00", "Hi"]
