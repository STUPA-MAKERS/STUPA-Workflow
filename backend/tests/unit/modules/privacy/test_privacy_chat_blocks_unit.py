"""Unit tests of the chat files and the sender blocks in the privacy module (no database).

The Auskunft names the files of the comments that the subject can see and lists the
sender blocks of the account or the address. The erasure of an account, of an address
and of an applicant request deletes the blocks.
"""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.modules.applications import service as applications_service
from app.modules.applications.models import Applicant, Application, Comment
from app.modules.auth.models import Principal
from app.modules.privacy import service as privacy_service
from app.modules.privacy.models import ErasureRequest
from app.modules.privacy.service import (
    AuskunftService,
    ErasureRequestService,
    PrincipalService,
    build_auskunft_workbook,
    erase_email,
)
from app.modules.tickets.models import TicketSenderBlock
from tests._support.privacy_fakes import fake_session, result

_AT = datetime(2026, 6, 16, 12, 0, tzinfo=UTC)


@pytest.fixture
def erased(monkeypatch: pytest.MonkeyPatch) -> list[dict[str, Any]]:
    seen: list[dict[str, Any]] = []

    async def _erase(_session: object, **kw: Any) -> None:
        seen.append(kw)

    async def _none(*_a: object, **_k: object) -> Any:
        return 0

    monkeypatch.setattr(privacy_service, "erase_sender_blocks", _erase)
    monkeypatch.setattr(privacy_service, "erase_candidacies", _none)
    monkeypatch.setattr(privacy_service, "erase_candidate_names", _none)
    monkeypatch.setattr(privacy_service, "sync_principal_memberships", _none)
    monkeypatch.setattr(privacy_service, "anonymize_reported_tickets", _none)
    return seen


async def test_the_account_erasure_deletes_the_blocks_of_the_account_and_its_email(
    erased: list[dict[str, Any]],
) -> None:
    principal = Principal(id=uuid.uuid4(), sub="kc", email="u@x.org", active=True)
    await PrincipalService(fake_session(gets=[principal])).erase(
        principal.id, actor="admin", commit=False
    )
    assert erased == [{"principal_id": principal.id, "email": "u@x.org"}]


async def test_the_address_erasure_deletes_its_block(erased: list[dict[str, Any]]) -> None:
    out = await erase_email(fake_session(), "a@x.org", actor="admin")
    assert out["applications"] == 0
    assert erased == [{"email": "a@x.org"}]


async def test_the_applicant_request_deletes_the_block_of_its_address(
    erased: list[dict[str, Any]], monkeypatch: pytest.MonkeyPatch
) -> None:
    async def _anonymize(*_a: object, **_k: object) -> None:
        return None

    monkeypatch.setattr(applications_service.ApplicationsService, "anonymize", _anonymize)
    request = ErasureRequest(
        id=uuid.uuid4(),
        subject_type="applicant",
        application_id=uuid.uuid4(),
        email="b@x.org",
        status="open",
    )
    await ErasureRequestService(fake_session(gets=[request])).execute(request.id, actor="dpo")
    assert erased == [{"email": "b@x.org"}]


async def test_the_export_names_the_comment_files_and_lists_the_blocks(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Name the comment files and list the blocks in the export.

    The last queued results are the versions, the application attachments and the
    tickets.
    """
    asked: list[dict[str, Any]] = []

    async def _blocks(_session: object, **kw: Any) -> list[TicketSenderBlock]:
        asked.append(kw)
        return [
            TicketSenderBlock(
                id=uuid.uuid4(),
                principal_id=kw["principal_id"],
                created_at=_AT,
                hit_count=2,
                last_hit_at=_AT,
            ),
            TicketSenderBlock(id=uuid.uuid4(), email="x@example.org", created_at=_AT, hit_count=0),
        ]

    async def _todos(_session: object, _ids: list[uuid.UUID]) -> dict[str, Any]:
        return {"cards": [], "comments": []}

    monkeypatch.setattr(privacy_service, "sender_blocks_of", _blocks)
    monkeypatch.setattr(privacy_service, "todo_export", _todos)
    app_id, type_id = uuid.uuid4(), uuid.uuid4()
    applicant = Applicant(application_id=app_id, email="x@example.org", name="X")
    application = Application(
        id=app_id, type_id=type_id, current_state_id=None, created_at=_AT, data={}
    )
    with_files = Comment(
        id=uuid.uuid4(),
        application_id=app_id,
        author_kind="applicant",
        body="",
        visibility="public",
        at=_AT,
    )
    plain = Comment(
        id=uuid.uuid4(),
        application_id=app_id,
        author_kind="principal",
        body="Antwort",
        visibility="public",
        at=_AT,
    )
    principal = Principal(id=uuid.uuid4(), sub="kc", email="x@example.org", active=True)
    db = fake_session(
        scalars=[
            result(applicant),
            result(application),
            result(),
            result(with_files, plain),
            result(),
            result(),
        ],
        execute=[
            result((type_id, {"de": "Typ"})),
            result((with_files.id, "a.pdf"), (with_files.id, "b.png")),
        ],
        scalar=[principal],
    )
    data = await AuskunftService(db).collect("x@example.org")
    assert [c["attachments"] for c in data["comments"]] == [["a.pdf", "b.png"], []]
    assert asked == [{"principal_id": principal.id, "email": "x@example.org"}]
    assert data["sender_blocks"] == [
        {"kind": "principal", "createdAt": _AT, "hitCount": 2, "lastHitAt": _AT},
        {"kind": "email", "createdAt": _AT, "hitCount": 0, "lastHitAt": None},
    ]


async def test_the_export_without_comments_asks_for_no_files(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    asked: list[dict[str, Any]] = []

    async def _blocks(_session: object, **kw: Any) -> list[TicketSenderBlock]:
        asked.append(kw)
        return []

    monkeypatch.setattr(privacy_service, "sender_blocks_of", _blocks)
    app_id, type_id = uuid.uuid4(), uuid.uuid4()
    applicant = Applicant(application_id=app_id, email="y@example.org", name="Y")
    application = Application(
        id=app_id, type_id=type_id, current_state_id=None, created_at=_AT, data={}
    )
    db = fake_session(
        scalars=[result(applicant), result(application)],
        execute=[result((type_id, {"de": "Typ"}))],
        scalar=[None],
    )
    data = await AuskunftService(db).collect("y@example.org")
    assert data["comments"] == []
    assert data["sender_blocks"] == []
    assert asked == [{"principal_id": None, "email": "y@example.org"}]


def test_the_workbook_names_the_comment_files_and_lists_the_blocks() -> None:
    raw = build_auskunft_workbook(
        email="x@example.org",
        applications=[],
        versions=[],
        principal=None,
        comments=[
            {
                "applicationId": "a1",
                "authorKind": "applicant",
                "visibility": "public",
                "body": "",
                "at": _AT,
                "attachments": ["a.pdf", "=b.png"],
            },
            {"applicationId": None, "body": None},
        ],
        sender_blocks=[
            {"kind": "principal", "createdAt": _AT, "hitCount": 2, "lastHitAt": _AT},
            {"kind": "email", "createdAt": _AT, "hitCount": 0, "lastHitAt": None},
        ],
    )
    book = load_workbook(BytesIO(raw))
    comments = [[c.value for c in row] for row in book["Kommentare"].iter_rows()]
    assert comments[0][-1] == "Anhänge"
    assert comments[1][-1] == "a.pdf; =b.png"
    assert comments[2][-1] is None
    blocks = [[c.value for c in row] for row in book["Ticket-Sperren"].iter_rows()]
    assert blocks[0] == ["Art", "gesperrt am", "abgefangene Meldungen", "letzter Versuch"]
    assert blocks[1] == ["Konto", "2026-06-16 12:00", 2, "2026-06-16 12:00"]
    assert blocks[2] == ["E-Mail-Adresse", "2026-06-16 12:00", 0, None]
