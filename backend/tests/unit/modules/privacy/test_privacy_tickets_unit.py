"""Unit tests of the ticket part of the privacy module (no database).

The erasure of an account anonymizes the tickets that it reported. The Art. 15
export lists the tickets that an address reported, as an account or as a bare
e-mail, in its own sheet.
"""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.modules.auth.models import Principal
from app.modules.privacy import service as privacy_service
from app.modules.privacy.service import (
    AuskunftService,
    PrincipalService,
    build_auskunft_workbook,
)
from app.modules.tickets.models import Ticket
from tests._support.privacy_fakes import fake_session, result
from tests.unit.modules.tickets.fakes import FakeFiles, ScriptedSession

_AT = datetime(2026, 6, 16, 12, 0, tzinfo=UTC)


def _ticket(**kw: Any) -> Ticket:
    base: dict[str, Any] = {
        "id": uuid.uuid4(),
        "number": 3,
        "kind": "bug",
        "title": "Kaputt",
        "body": "b",
        "status": "done",
        "reporter_principal_id": None,
        "reporter_email": None,
        "context": {"url": "u"},
        "anonymized_at": None,
        "created_at": _AT,
    }
    base.update(kw)
    return Ticket(**base)


@pytest.fixture(autouse=True)
def _quiet(monkeypatch: pytest.MonkeyPatch) -> None:
    async def _none(*_a: object, **_k: object) -> list[object]:
        return []

    async def _synced(*_a: object, **_k: object) -> bool:
        return True

    monkeypatch.setattr(privacy_service, "erase_candidacies", _none)
    monkeypatch.setattr(privacy_service, "erase_candidate_names", _none)
    monkeypatch.setattr(privacy_service, "sync_principal_memberships", _synced)


async def test_the_erasure_anonymizes_the_reported_tickets() -> None:
    principal = Principal(id=uuid.uuid4(), sub="kc", email="u@x.org", active=True)
    ticket = _ticket(reporter_principal_id=principal.id)
    # The first result is the comment lookup of ``lock_personal_todos``, which runs
    # before the first audit entry of the erasure.
    db = fake_session(gets=[principal], scalars=[result(), result(ticket)])
    files = FakeFiles()
    await PrincipalService(db).erase(
        principal.id,
        actor="admin",
        commit=False,
        files=files,  # type: ignore[arg-type]
    )
    assert ticket.reporter_principal_id is None
    assert ticket.context == {}
    assert ticket.anonymized_at is not None
    assert files.deleted == [(ticket.id, "admin")]
    assert files.locked == [("todo_comment", [])]


async def test_the_erasure_without_tickets_builds_its_own_files_service() -> None:
    principal = Principal(id=uuid.uuid4(), sub="kc", email="u@x.org", active=True)
    db = fake_session(gets=[principal])
    out = await PrincipalService(db).erase(principal.id, actor="admin", commit=False)
    assert out.active is False


async def test_the_export_lists_the_tickets_of_the_address() -> None:
    mine = _ticket(reporter_email="x@example.org")
    db = fake_session(scalars=[result(), result(mine)], scalar=[None])
    data = await AuskunftService(db).collect("x@example.org")
    assert data["tickets"] == [
        {
            "number": "T-0003",
            "title": "Kaputt",
            "status": "done",
            "createdAt": _AT,
            "body": "b",
            "context": {"url": "u"},
        }
    ]


async def test_the_export_lists_the_attachments_of_the_reported_tickets() -> None:
    from app.modules.files.models import Attachment

    mine = _ticket(reporter_email="x@example.org", number=7)
    shot = Attachment(
        id=uuid.uuid4(),
        ticket_id=mine.id,
        filename="screen.png",
        mime="image/png",
        size=12,
        created_at=_AT,
    )
    db = fake_session(scalars=[result(), result(mine), result(shot)], scalar=[None])
    data = await AuskunftService(db).collect("x@example.org")
    assert data["attachments"] == [
        {
            "ticketNumber": "T-0007",
            "filename": "screen.png",
            "mime": "image/png",
            "size": 12,
            "createdAt": _AT,
        }
    ]


async def test_the_export_lists_the_tickets_of_the_account() -> None:
    principal = Principal(id=uuid.uuid4(), sub="kc", email="x@example.org", active=True)
    ticket = _ticket(reporter_principal_id=principal.id, number=12)
    db = fake_session(scalars=[result(), result(ticket)], scalar=[principal])
    data = await AuskunftService(db).collect("x@example.org")
    assert [t["number"] for t in data["tickets"]] == ["T-0012"]


def test_the_workbook_has_a_ticket_sheet() -> None:
    raw = build_auskunft_workbook(
        email="x@example.org",
        applications=[],
        versions=[],
        principal=None,
        tickets=[
            {
                "number": "T-0001",
                "title": "=HYPERLINK()",
                "status": "new",
                "createdAt": _AT,
                "body": "Der Knopf geht nicht.",
                "context": {"url": "/x", "userAgent": "UA"},
            },
            {"number": None, "title": None, "status": None, "createdAt": None},
        ],
        attachments=[
            {
                "ticketNumber": "T-0001",
                "filename": "screen.png",
                "mime": "image/png",
                "size": 12,
                "createdAt": _AT,
            }
        ],
    )
    book = load_workbook(BytesIO(raw))
    rows = [[c.value for c in row] for row in book["Tickets"].iter_rows()]
    assert rows[0] == ["Nummer", "Titel", "Status", "Gemeldet", "Text", "Technischer Kontext"]
    assert rows[1][0] == "T-0001"
    assert not str(rows[1][1]).startswith("=")
    assert rows[1][4] == "Der Knopf geht nicht."
    assert rows[1][5] == '{"url": "/x", "userAgent": "UA"}'
    assert rows[2] == [None] * 6
    files = [[c.value for c in row] for row in book["Anhänge"].iter_rows()]
    assert files[0][0] == "Antrag / Ticket"
    assert files[1][:2] == ["T-0001", "screen.png"]


def test_the_erase_route_passes_the_object_storage(monkeypatch: pytest.MonkeyPatch) -> None:
    """The direct erasure removes the ticket attachment objects too."""
    from collections.abc import Iterator

    from fastapi.testclient import TestClient

    from app.db import get_session
    from app.deps import get_current_principal
    from app.main import create_app
    from app.modules.auth.principal import Principal as SessionPrincipal

    seen: dict[str, Any] = {}

    async def _erase(_self: object, principal_id: uuid.UUID, **kw: Any) -> None:
        seen.update(kw, principal_id=principal_id)

    monkeypatch.setattr(PrincipalService, "erase", _erase)
    app = create_app()
    storage = object()
    app.state.object_storage = storage

    def _session() -> Iterator[object]:
        yield object()

    app.dependency_overrides[get_session] = _session
    app.dependency_overrides[get_current_principal] = lambda: SessionPrincipal(
        sub="admin", permissions={"privacy.manage"}
    )
    target = uuid.uuid4()
    with TestClient(app) as client:
        app.state.object_storage = storage
        resp = client.post(f"/api/admin/privacy/principals/{target}/erase")
    assert resp.status_code == 204
    assert seen["principal_id"] == target
    assert seen["actor"] == "admin"
    assert seen["files"].storage is storage


def _sql(stmt: object) -> str:
    from sqlalchemy.dialects import postgresql

    return str(stmt.compile(dialect=postgresql.dialect()))  # type: ignore[attr-defined]


async def test_the_reported_tickets_without_a_reporter_are_none() -> None:
    session = ScriptedSession()
    count = await privacy_service.anonymize_reported_tickets(
        session,  # type: ignore[arg-type]
        actor="admin",
        email=None,
    )
    assert count == 0
    assert session.statements == []


async def test_the_reported_tickets_match_the_bare_e_mail() -> None:
    ticket = _ticket(reporter_email="x@example.org")
    session = ScriptedSession(scalars=[[ticket]])
    files = FakeFiles()
    count = await privacy_service.anonymize_reported_tickets(
        session,  # type: ignore[arg-type]
        actor="admin",
        files=files,  # type: ignore[arg-type]
        email="x@example.org",
    )
    assert count == 1
    query = _sql(session.statements[0])
    assert "reporter_email" in query
    assert "reporter_principal_id =" not in query
    assert "anonymized_at IS NULL" in query
    # The erasure locks the tickets like add_note and the retention cron, so a note
    # cannot bind files while the ticket is anonymized.
    assert query.endswith("ORDER BY ticket.id FOR UPDATE")
    assert ticket.reporter_email is None
    assert ticket.anonymized_at is not None
    assert files.deleted == [(ticket.id, "admin")]


async def test_the_reported_tickets_match_the_account_and_its_e_mail() -> None:
    session = ScriptedSession(scalars=[[]])
    count = await privacy_service.anonymize_reported_tickets(
        session,  # type: ignore[arg-type]
        actor="admin",
        principal_id=uuid.uuid4(),
        email="u@x.org",
    )
    assert count == 0
    query = _sql(session.statements[0])
    assert "reporter_principal_id =" in query
    assert "reporter_email =" in query


async def test_the_erasure_passes_the_e_mail_from_before_the_erasure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seen: dict[str, Any] = {}

    async def _tickets(_session: object, **kw: Any) -> int:
        seen.update(kw)
        return 0

    monkeypatch.setattr(privacy_service, "anonymize_reported_tickets", _tickets)
    principal = Principal(id=uuid.uuid4(), sub="kc", email="u@x.org", active=True)
    db = fake_session(gets=[principal])
    await PrincipalService(db).erase(principal.id, actor="admin", commit=False)
    assert seen["principal_id"] == principal.id
    assert seen["email"] == "u@x.org"
    assert principal.email is None


async def test_the_applicant_erasure_anonymizes_the_tickets_of_its_e_mail(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.modules.applications import service as applications_service
    from app.modules.privacy.models import ErasureRequest
    from app.modules.privacy.service import ErasureRequestService

    seen: dict[str, Any] = {}

    async def _anonymize(*_a: object, **_k: object) -> None:
        return None

    async def _tickets(_session: object, **kw: Any) -> int:
        seen.update(kw)
        return 1

    monkeypatch.setattr(applications_service.ApplicationsService, "anonymize", _anonymize)
    monkeypatch.setattr(privacy_service, "anonymize_reported_tickets", _tickets)
    request = ErasureRequest(
        id=uuid.uuid4(),
        subject_type="applicant",
        application_id=uuid.uuid4(),
        email="a@example.org",
        status="open",
    )
    db = fake_session(gets=[request])
    out = await ErasureRequestService(db).execute(request.id, actor="dpo")
    assert out.status == "executed"
    assert seen == {"actor": "dpo", "files": None, "email": "a@example.org"}
